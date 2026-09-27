import frappe
import openpyxl
import io

@frappe.whitelist()
def process_kpi_upload(docname):
    doc = frappe.get_doc("Rider KPI Upload", docname)

    if not doc.attach_file:
        frappe.throw("No file attached")

    file_doc = frappe.get_doc("File", {"file_url": doc.attach_file})
    content = file_doc.get_content()

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active

    headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]

    doc.set("riders_data", [])

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        row_data = dict(zip(headers, row))

        doc.append("riders_data", {
            "rider_id": str(row_data.get("rider_id") or ""),
            "rider_name": str(row_data.get("Name") or ""),
            "contract_name": str(row_data.get("contract_name") or ""),
            "planned_hours": float(row_data.get("Planned Hours") or 0),
            "actual_working_hours": float(row_data.get("Actual Working Hours") or 0),
            "accepted_count": int(row_data.get("accepted_count") or 0),
            "notified_count": int(row_data.get("notified_count") or 0),
            "manual_undispatch_count": int(row_data.get("manual_undispatch_count") or 0),
            "lost_hours": float(row_data.get("Lost Hours") or 0),
            "no_show_shifts": int(row_data.get("no_show_shifts") or 0),
            "break_hours": float(row_data.get("Break Hours") or 0),
            "total_orders": int(row_data.get("Total Orders") or 0),
            "total_deliveries": int(row_data.get("Total Deliveries") or 0),
        })

    doc.save()
    frappe.db.commit()
    return "Success"
